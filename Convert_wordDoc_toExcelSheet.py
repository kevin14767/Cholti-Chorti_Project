#this python script is intended to convert "The Religious Sections" of the Moran Manuscript into a database file
#Excel sheet format will be a answer and response format, naturally the manuscript follows these format as well

#Formatting:
#Line Number 1-530
#Manuscript entry Page.59 to 80
#Line entry in the actual manuscript page. i.e entry number 1 entry number 2 however sometimes they are in the same line entry or in
#multiple lines for one entry. 
#4 columns: standardized version of the original ch'olti, morphemic gloss, a literal english translation, a flowing english translation
from docx import Document
from openpyxl import load_workbook

doc = Document("liturgy30June06.docx")

workbook = load_workbook("LiturgyAndOriginalTranscription.xlsx")
sheet = workbook.active

for paragraphs in doc.paragraphs:
    lines = paragraphs.text.split('\n')
    print(lines)